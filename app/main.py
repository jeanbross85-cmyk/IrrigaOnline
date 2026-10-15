from pathlib import Path
from datetime import datetime
from io import BytesIO
import json, re, shutil, secrets, uuid
from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import (
    HTMLResponse,
    RedirectResponse,
    StreamingResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import select, desc
from sqlalchemy.exc import IntegrityError
from .config import APP_NAME, KC_ESTADIOS, SECRET_KEY, COOKIE_SECURE
from .db import SessionLocal, init_db
from .models import (
    DailyRecord,
    User,
    AuditLog,
    PendingImport,
    DrainageRecord,
)
from .security import current_user, require_user, verify_password, ensure_csrf
from .importer import read_meteorological_xlsx
from .calculations import daily_summary, irrigation, coverage_status
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
DATA_DIR = PROJECT_DIR / "data"
PENDING_DIR = PROJECT_DIR / "instance" / "pending"
DATA_DIR.mkdir(exist_ok=True); PENDING_DIR.mkdir(parents=True, exist_ok=True)
app = FastAPI(title=APP_NAME)
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, same_site="lax", https_only=COOKIE_SECURE, max_age=60*60*8)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
@app.on_event("startup")
def startup(): init_db()
def template(name, request, **kwargs):
    if "csrf_token" not in request.session: request.session["csrf_token"] = secrets.token_urlsafe(32)
    return templates.TemplateResponse(name, {"request": request, "app_name": APP_NAME, "user": current_user(request), "csrf_token": request.session["csrf_token"], **kwargs})
def auth_redirect(request):
    if not current_user(request): return RedirectResponse("/login", status_code=303)
    return None
@app.get("/login", response_class=HTMLResponse)
def login_get(request: Request):
    if current_user(request): return RedirectResponse("/", status_code=303)
    return template("login.html", request, error=None)
@app.post("/login", response_class=HTMLResponse)
def login_post(request: Request, username: str = Form(...), password: str = Form(...), csrf_token: str = Form(...)):
    ensure_csrf(request, csrf_token)
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == username.strip()))
        if not user or not user.active or not verify_password(password, user.password_hash):
            return template("login.html", request, error="Usuário ou senha inválidos.")
        request.session.clear(); request.session["user_id"] = user.id; request.session["csrf_token"] = secrets.token_urlsafe(32)
    return RedirectResponse("/", status_code=303)
@app.post("/logout")
def logout(request: Request, csrf_token: str = Form(...)):
    ensure_csrf(request, csrf_token); request.session.clear()
    return RedirectResponse("/login", status_code=303)

@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request):
    redirect = auth_redirect(request)
    if redirect:
        return redirect

    with SessionLocal() as db:
        records = list(
            db.scalars(
                select(DailyRecord)
                .order_by(desc(DailyRecord.data))
                .limit(30)
            )
        )

        drainage_records = list(
            db.scalars(
                select(DrainageRecord)
                .order_by(desc(DrainageRecord.data))
                .limit(30)
            )
        )

    return template(
        "dashboard.html",
        request,
        records=records,
        latest=records[0] if records else None,
        drainage_records=drainage_records,
        latest_drainage=(
            drainage_records[0] if drainage_records else None
        ),
    )
@app.get("/historico", response_class=HTMLResponse)
def historico(request: Request):
    if auth_redirect(request): return auth_redirect(request)
    with SessionLocal() as db: records = list(db.scalars(select(DailyRecord).order_by(desc(DailyRecord.data))))
    return template("historico.html", request, records=records)

@app.get("/historico/exportar.xlsx")
def exportar_historico_excel(request: Request):
    redirect = auth_redirect(request)
    if redirect:
        return redirect

    # Consultar os dados existentes sem modificá-los.
    with SessionLocal() as db:
        records = list(
            db.scalars(
                select(DailyRecord).order_by(DailyRecord.data)
            )
        )
        drainage_records = list(
            db.scalars(
                select(DrainageRecord).order_by(DrainageRecord.data)
            )
        )

        # Criar a pasta de trabalho.
        wb = Workbook()
        ws = wb.active
        ws.title = "Historico Irrigacao"

        # Nomes amigáveis para as colunas conhecidas.
        nomes = {
            "data": "Data",
            "estadio": "Estádio",
            "kc": "Kc",
            "eto_mm": "ETo (mm)",
            "etc_com_kl_mm": "ETc com Kl (mm)",
            "etc_sem_kl_mm": "ETc sem Kl (mm)",
            "temperatura_media_c": "Temperatura média (°C)",
            "chuva_mm": "Precipitação (mm)",
            "rn_mj_m2_dia": "Saldo de radiação (MJ/m²/dia)",
            "lamina_bruta_com_kl_mm": "Lâmina bruta com Kl (mm)",
            "lamina_bruta_sem_kl_mm": "Lâmina bruta sem Kl (mm)",
            "volume_planta_com_kl_l": "Volume por planta com Kl (L)",
            "volume_planta_sem_kl_l": "Volume por planta sem Kl (L)",
            "tempo_irrigacao_com_kl_min": "Tempo com Kl (min)",
            "tempo_irrigacao_sem_kl_min": "Tempo sem Kl (min)",
            "cobertura_dados": "Cobertura dos dados",
            "status_dados": "Status dos dados",
        }

        # Usar todas as colunas reais do modelo.
        colunas = list(DailyRecord.__table__.columns)

        ws.append([
            nomes.get(col.name, col.name)
            for col in colunas
        ])

        for record in records:
            linha = []
            for col in colunas:
                valor = getattr(record, col.name)

                # Converter enums em texto, quando aplicável.
                if hasattr(valor, "value"):
                    valor = valor.value

                linha.append(valor)

            ws.append(linha)

        # Segunda aba: medições de drenagem.
        wd = wb.create_sheet("Historico Drenagem")

        colunas_drenagem = list(
            DrainageRecord.__table__.columns
        )

        nomes_drenagem = {
            "data": "Data",
            "leitura_percentual": "Leitura do sensor (%)",
            "umidade_volumetrica": "Umidade volumétrica (m³/m³)",
        }

        wd.append([
            nomes_drenagem.get(col.name, col.name)
            for col in colunas_drenagem
        ])

        for record in drainage_records:
            linha = []
            for col in colunas_drenagem:
                valor = getattr(record, col.name)

                if hasattr(valor, "value"):
                    valor = valor.value

                linha.append(valor)

            wd.append(linha)

    # Formatar os cabeçalhos das duas planilhas.
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="245A45",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sheet.row_dimensions[1].height = 32

        for column_cells in sheet.columns:
            largura = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            letra = column_cells[0].column_letter
            sheet.column_dimensions[letra].width = min(
                max(largura + 2, 12), 35
            )

    # Gerar o arquivo em memória.
    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    wb.close()

    return StreamingResponse(
        arquivo,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="historico_irrigao.xlsx"'
        },
    )
def records_as_json(records):
    result = []

    for r in records:
        result.append({
            "data": r.data.isoformat(),
            "eto": r.eto_mm,

            # Campos antigos: mantidos para compatibilidade.
            "etc": r.etc_mm,
            "lamina": r.lamina_bruta_mm,
            "volume_planta": r.volume_planta_l,
            "tempo": r.tempo_irrigacao_min,

            # Cenários com e sem Kl.
            "etc_com_kl_mm": r.etc_com_kl_mm,
            "etc_sem_kl_mm": r.etc_sem_kl_mm,
            "lamina_bruta_com_kl_mm": r.lamina_bruta_com_kl_mm,
            "lamina_bruta_sem_kl_mm": r.lamina_bruta_sem_kl_mm,
            "volume_planta_com_kl_l": r.volume_planta_com_kl_l,
            "volume_planta_sem_kl_l": r.volume_planta_sem_kl_l,
            "tempo_irrigacao_com_kl_min": r.tempo_irrigacao_com_kl_min,
            "tempo_irrigacao_sem_kl_min": r.tempo_irrigacao_sem_kl_min,

            # Outras variáveis do histórico.
            "temperatura": r.temperatura_media_c,
            "chuva": r.chuva_mm,
            "cobertura": r.cobertura_dados,
            "status": r.status_dados,
        })

    return result
@app.get("/api/historico")
def api_historico(request: Request):
    if auth_redirect(request): return auth_redirect(request)
    with SessionLocal() as db: records=list(db.scalars(select(DailyRecord).order_by(DailyRecord.data)))
    return records_as_json(records)
@app.get("/api/drenagem")
def api_drenagem(request: Request):
    redirect = auth_redirect(request)
    if redirect:
        return redirect

    with SessionLocal() as db:
        registros = list(
            db.scalars(
                select(DrainageRecord).order_by(
                    DrainageRecord.data,
                    DrainageRecord.criado_em,
                    DrainageRecord.id,
                )
            ).all()
        )

        return [
            {
                "data": r.data.isoformat(),
                "leitura_percentual": r.leitura_percentual,
                "umidade_volumetrica": r.umidade_volumetrica,
            }
            for r in registros
        ]
@app.get("/importar", response_class=HTMLResponse)
def importar_get(request: Request):
    if auth_redirect(request): return auth_redirect(request)
    return template("importar.html", request, message=None, error=None, conflicts=None, pending_id=None)
def parse_filename(filename):
    return bool(re.fullmatch(r"Dados_meteorologicos_\d{4}-\d{2}-\d{2}\.xlsx", filename, flags=re.IGNORECASE))
def build_day_values(day, day_df, estadio, filename, user_id):
    s = daily_summary(day_df)
    kc = KC_ESTADIOS[estadio]
    irr = irrigation(s["eto"], kc, s["chuva"])

    return {
        "data": day,
        "arquivo_estacao": filename,
        "estadio": estadio,
        "kc": kc,
        "eto_mm": s["eto"],

        # Campos antigos: preservam o cenário com Kl.
        "etc_mm": irr["etc_com_kl_mm"],
        "chuva_efetiva_mm": irr["chuva_efetiva"],
        "necessidade_liquida_mm": irr["necessidade_liquida"],
        "lamina_bruta_mm": irr["lamina_bruta"],
        "volume_planta_l": irr["volume_planta"],
        "tempo_irrigacao_min": irr["tempo_irrigacao"],

        # ETc com e sem Kl.
        "etc_com_kl_mm": irr["etc_com_kl_mm"],
        "etc_sem_kl_mm": irr["etc_sem_kl_mm"],

        # Lâmina bruta com e sem Kl.
        "lamina_bruta_com_kl_mm": irr["lamina_bruta_com_kl_mm"],
        "lamina_bruta_sem_kl_mm": irr["lamina_bruta_sem_kl_mm"],

        # Volume por planta com e sem Kl.
        "volume_planta_com_kl_l": irr["volume_planta_com_kl_l"],
        "volume_planta_sem_kl_l": irr["volume_planta_sem_kl_l"],

        # Tempo de irrigação com e sem Kl.
        "tempo_irrigacao_com_kl_min": irr["tempo_irrigacao_com_kl_min"],
        "tempo_irrigacao_sem_kl_min": irr["tempo_irrigacao_sem_kl_min"],

        # Dados meteorológicos.
        "temperatura_media_c": s["temperatura_media"],
        "temperatura_max_c": s["temperatura_max"],
        "temperatura_min_c": s["temperatura_min"],
        "rh_max": s["rh_max"],
        "rh_min": s["rh_min"],
        "vento_ms": s["vento"],
        "chuva_mm": s["chuva"],
        "rn_mj_m2_dia": s["rn"],

        # Qualidade dos dados e auditoria.
        "cobertura_dados": s["cobertura"],
        "status_dados": coverage_status(s["cobertura"]),
        "atualizado_por_id": user_id,
    }
def snapshot(record):
    if not record: return None
    return {c.name:(getattr(record,c.name).isoformat() if hasattr(getattr(record,c.name),"isoformat") else getattr(record,c.name)) for c in DailyRecord.__table__.columns}
def apply_import(path, filename, estadio, user):
    df = read_meteorological_xlsx(path)

    if df.empty:
        raise ValueError(
            "O arquivo não contém registros com data/hora válida."
        )

    days = []

    with SessionLocal() as db:
        try:
            for day, day_df in df.groupby("DATA"):
                values = build_day_values(
                    day, day_df, estadio, filename, user["id"]
                )

                rec = db.scalar(
                    select(DailyRecord)
                    .where(DailyRecord.data == day)
                    .with_for_update()
                )

                before = snapshot(rec)
                action = "ATUALIZACAO" if rec else "CRIACAO"

                if rec is None:
                    rec = DailyRecord(**values)
                    db.add(rec)
                else:
                    for key, value in values.items():
                        setattr(rec, key, value)

                db.flush()

                db.add(
                    AuditLog(
                        data_registro=day,
                        acao=action,
                        usuario_id=user["id"],
                        usuario_nome=user["username"],
                        arquivo=filename,
                        antes_json=(
                            json.dumps(
                                before,
                                ensure_ascii=False,
                                default=str,
                            )
                            if before
                            else None
                        ),
                        depois_json=json.dumps(
                            snapshot(rec),
                            ensure_ascii=False,
                            default=str,
                        ),
                    )
                )

                days.append(day)

            db.commit()

        except IntegrityError as exc:
            db.rollback()
            raise ValueError(
                "Outra pessoa atualizou uma das datas ao mesmo tempo. "
                "O banco protegeu o histórico contra duplicação. "
                "Atualize a página e repita a importação."
            ) from exc

    return len(days), days

def concluir_importacao_pendente(pending_id, user_id):
    """
    Salva a medição de drenagem, se informada, e remove a importação
    pendente após a importação meteorológica ter sido concluída.
    """
    with SessionLocal() as db:
        pending = db.get(PendingImport, pending_id)

        if not pending or pending.user_id != user_id:
            raise ValueError(
                "A importação pendente não foi encontrada."
            )

        leitura = pending.drenagem_leitura_percentual

        if leitura is not None:
            u = float(leitura)
            theta = 0.0008 * (u ** 2) - 0.0155 * u + 0.1203

            db.add(
                DrainageRecord(
                    data=pending.drenagem_data,
                    leitura_percentual=u,
                    umidade_volumetrica=theta,
                    arquivo_origem=pending.filename,
                    usuario_id=user_id,
                )
            )

        db.delete(pending)
        db.commit()


@app.post("/importar/processar", response_class=HTMLResponse)
def processar_importacao(
    request: Request,
    pending_id: str = Form(...),
    drenagem_data: str = Form(...),
    drenagem_leitura_percentual: str = Form(""),
    csrf_token: str = Form(...),
):
    redirect = auth_redirect(request)
    if redirect:
        return redirect

    ensure_csrf(request, csrf_token)
    user = current_user(request)

    # Valida e interpreta a data escolhida.
    try:
        data_medicao = datetime.strptime(
            drenagem_data, "%Y-%m-%d"
        ).date()
    except (TypeError, ValueError):
        return template(
            "importar.html", request,
            message=None,
            error="Informe uma data válida para a medição.",
            conflicts=None,
            pending_id=pending_id,
            show_drainage_form=True,
            suggested_date=drenagem_data,
        )

    # A leitura é opcional, mas, quando informada, deve estar entre 0 e 100.
    leitura = None
    if drenagem_leitura_percentual.strip():
        try:
            leitura = float(
                drenagem_leitura_percentual.strip().replace(",", ".")
            )
        except ValueError:
            leitura = None
            return template(
                "importar.html", request,
                message=None,
                error="Informe uma leitura numérica válida, em porcentagem.",
                conflicts=None,
                pending_id=pending_id,
                show_drainage_form=True,
                suggested_date=drenagem_data,
                suggested_reading=drenagem_leitura_percentual,

            )

        if not 0 <= leitura <= 100:
            return template(
                "importar.html", request,
                message=None,
                error="A leitura deve estar entre 0 e 100%.",
                conflicts=None,
                pending_id=pending_id,
                show_drainage_form=True,
                suggested_date=drenagem_data,
                suggested_reading=drenagem_leitura_percentual,
            )


    with SessionLocal() as db:
        pending = db.get(PendingImport, pending_id)

        if not pending or pending.user_id != user["id"]:
            raise HTTPException(
                404,
                "Importação pendente não encontrada ou não pertence a este usuário."
            )

        # Verifica a duplicidade antes de guardar a leitura
        # ou importar os dados meteorológicos.
        if leitura is not None:
            drenagem_existente = db.scalar(
                select(DrainageRecord.id)
                .where(DrainageRecord.data == data_medicao)
                .limit(1)
            )
            if drenagem_existente is not None:
                return template(
                    "importar.html",
                    request,
                    message=None,
                    error=(
                        "Já existe uma medição de drenagem para "
                        f"{data_medicao.strftime('%d/%m/%Y')}. "
                        "Escolha outra data para registrar uma nova leitura."
                    ),
                    conflicts=None,
                    pending_id=pending_id,
                    show_drainage_form=True,
                    suggested_date=drenagem_data,
                    suggested_reading=leitura,
                )

        # Só guarda os dados após validar a data.
        pending.drenagem_data = data_medicao
        pending.drenagem_leitura_percentual = leitura

        filename = pending.filename
        estadio = pending.estadio
        file_bytes = bytes(pending.file_content)

        db.commit()

    path = PENDING_DIR / f"process-{uuid.uuid4()}.xlsx"
    path.write_bytes(file_bytes)

    try:
        df = read_meteorological_xlsx(path)
        dates = sorted(set(df["DATA"].tolist()))

        with SessionLocal() as db:
            existing = sorted(
                db.scalars(
                    select(DailyRecord.data).where(
                        DailyRecord.data.in_(dates)
                    )
                ).all()
            )

            if existing:
                pending = db.get(PendingImport, pending_id)
                if pending:
                    pending.conflicts_json = json.dumps(
                        [d.isoformat() for d in existing]
                    )
                    db.commit()

                return template(
                    "importar.html", request,
                    message=None,
                    error=None,
                    conflicts=[
                        d.strftime("%d/%m/%Y") for d in existing
                    ],
                    pending_id=pending_id,
                    show_drainage_form=False,
                )

        count, days = apply_import(path, filename, estadio, user)
        concluir_importacao_pendente(pending_id, user["id"])

        mensagem = (
            f"Importação concluída: {count} dia(s) gravado(s) "
            "no histórico compartilhado."
        )
        if leitura is not None:
            mensagem += (
                " Medição de drenagem salva para "
                f"{data_medicao.strftime('%d/%m/%Y')}."
            )

        return template(
            "importar.html", request,
            message=mensagem,
            error=None,
            conflicts=None,
            pending_id=None,
            show_drainage_form=False,
        )

    except Exception as exc:
        return template(
            "importar.html", request,
            message=None,
            error=f"Não foi possível concluir a importação: {exc}",
            conflicts=None,
            pending_id=pending_id,
            show_drainage_form=False,
        )

    finally:
        path.unlink(missing_ok=True)

@app.post("/importar", response_class=HTMLResponse)
async def importar_post(
    request: Request,
    arquivo: UploadFile = File(...),
    estadio: str = Form("Vegetativo"),
    csrf_token: str = Form(...),
):
    redirect = auth_redirect(request)
    if redirect:
        return redirect

    ensure_csrf(request, csrf_token)
    user = current_user(request)

    if not arquivo.filename or not arquivo.filename.lower().endswith(".xlsx"):
        return template(
            "importar.html", request,
            message=None,
            error="Selecione um arquivo Excel .xlsx.",
            conflicts=None, pending_id=None,
            show_drainage_form=False,
        )

    safe_name = Path(arquivo.filename).name

    if not parse_filename(safe_name):
        return template(
            "importar.html", request,
            message=None,
            error=(
                "O nome deve seguir o padrão "
                "Dados_meteorologicos_aaaa-mm-dd.xlsx, "
                "por exemplo Dados_meteorologicos_2026-10-09.xlsx."
            ),
            conflicts=None, pending_id=None,
            show_drainage_form=False,
        )

    if estadio not in KC_ESTADIOS:
        raise HTTPException(400, "Estágio inválido.")

    file_bytes = await arquivo.read()

    if len(file_bytes) > 25 * 1024 * 1024:
        return template(
            "importar.html", request,
            message=None,
            error="O arquivo excede o limite de 25 MB.",
            conflicts=None, pending_id=None,
            show_drainage_form=False,
        )

    pending_path = PENDING_DIR / f"{uuid.uuid4()}.xlsx"
    pending_path.write_bytes(file_bytes)

    try:
        df = read_meteorological_xlsx(pending_path)

        if df.empty:
            raise ValueError(
                "O arquivo não contém registros com data/hora válida."
            )

        # Sugere a data mais recente encontrada no Excel.
        suggested_date = max(df["DATA"].tolist())
        pending_id = str(uuid.uuid4())

        with SessionLocal() as db:
            pending = PendingImport(
                id=pending_id,
                user_id=user["id"],
                filename=safe_name,
                file_content=file_bytes,
                estadio=estadio,
                conflicts_json=json.dumps([]),
                drenagem_data=suggested_date,
                drenagem_leitura_percentual=None,
            )
            db.add(pending)
            db.commit()

        return template(
            "importar.html", request,
            message=None,
            error=None,
            conflicts=None,
            pending_id=pending_id,
            show_drainage_form=True,
            suggested_date=suggested_date.isoformat(),
        )

    except Exception as exc:
        return template(
            "importar.html", request,
            message=None,
            error=f"Não foi possível ler o arquivo: {exc}",
            conflicts=None,
            pending_id=None,
            show_drainage_form=False,
        )

    finally:
        pending_path.unlink(missing_ok=True)


@app.post("/importar/confirmar", response_class=HTMLResponse)
def confirmar_importacao(
    request: Request,
    pending_id: str = Form(...),
    csrf_token: str = Form(...),
):
    redirect = auth_redirect(request)
    if redirect:
        return redirect

    ensure_csrf(request, csrf_token)
    user = current_user(request)

    with SessionLocal() as db:
        pending = db.get(PendingImport, pending_id)

        if not pending or pending.user_id != user["id"]:
            raise HTTPException(
                404,
                "Importação pendente não encontrada ou não pertence a este usuário."
            )

        filename = pending.filename
        estadio = pending.estadio
        file_bytes = bytes(pending.file_content)
        previously_shown = set(
            json.loads(pending.conflicts_json or "[]")
        )

        drenagem_data = pending.drenagem_data
        drenagem_leitura = pending.drenagem_leitura_percentual

        if drenagem_leitura is not None and drenagem_data is not None:
            drenagem_existente = db.scalar(
                select(DrainageRecord.id)
                .where(DrainageRecord.data == drenagem_data)
                .limit(1)
            )

            if drenagem_existente is not None:
                return template(
                    "importar.html",
                    request,
                    message=None,
                    error=(
                        "Já existe uma medição de drenagem para "
                        f"{drenagem_data.strftime('%d/%m/%Y')}. "
                        "Escolha outra data para registrar a leitura."
                    ),
                    conflicts=None,
                    pending_id=pending_id,
                    show_drainage_form=True,
                    suggested_reading=drenagem_leitura,
                    suggested_date=drenagem_data.isoformat(),
                )

    path = PENDING_DIR / f"confirm-{uuid.uuid4()}.xlsx"
    path.write_bytes(file_bytes)

    try:
        df = read_meteorological_xlsx(path)
        dates = sorted(set(df["DATA"].tolist()))

        with SessionLocal() as db:
            now_existing = list(
                db.scalars(
                    select(DailyRecord.data).where(
                        DailyRecord.data.in_(dates)
                    )
                )
            )

            newly_conflicting = [
                d for d in now_existing
                if d.isoformat() not in previously_shown
            ]

            if newly_conflicting:
                with SessionLocal() as update_db:
                    staged = update_db.get(PendingImport, pending_id)
                    if staged:
                        staged.conflicts_json = json.dumps(
                            [d.isoformat() for d in sorted(now_existing)]
                        )
                        update_db.commit()

                return template(
                    "importar.html",
                    request,
                    message=None,
                    error=None,
                    conflicts=[
                        d.strftime("%d/%m/%Y")
                        for d in sorted(now_existing)
                    ],
                    pending_id=pending_id,
                )

        count, days = apply_import(path, filename, estadio, user)
        concluir_importacao_pendente(pending_id, user["id"])

        return template(
            "importar.html",
            request,
            message=(
                f"Atualização confirmada: {count} dia(s) inserido(s) "
                "ou substituído(s). O histórico de alterações foi registrado."
            ),
            error=None,
            conflicts=None,
            pending_id=None,
        )

    except Exception as exc:
        return template(
            "importar.html",
            request,
            message=None,
            error=f"Não foi possível confirmar a atualização: {exc}",
            conflicts=None,
            pending_id=pending_id,
        )

    finally:
        path.unlink(missing_ok=True)
@app.post("/importar/cancelar", response_class=HTMLResponse)
def cancelar_importacao(request: Request,pending_id: str=Form(...),csrf_token: str=Form(...)):
    redirect=auth_redirect(request)
    if redirect:return redirect
    ensure_csrf(request,csrf_token); user=current_user(request)
    with SessionLocal() as db:
        pending=db.get(PendingImport,pending_id)
        if pending and pending.user_id==user["id"]:
            db.delete(pending); db.commit()
    return template("importar.html",request,message="Importação cancelada; o histórico oficial não foi alterado.",error=None,conflicts=None,pending_id=None)
@app.get("/configuracoes", response_class=HTMLResponse)
def configuracoes(request: Request):
    redirect=auth_redirect(request)
    if redirect:return redirect
    return template("configuracoes.html",request,kc=KC_ESTADIOS)
@app.get("/auditoria", response_class=HTMLResponse)
def auditoria(request: Request):
    redirect=auth_redirect(request)
    if redirect:return redirect
    with SessionLocal() as db: logs=list(db.scalars(select(AuditLog).order_by(desc(AuditLog.alterado_em)).limit(300)))
    return template("auditoria.html",request,logs=logs)
